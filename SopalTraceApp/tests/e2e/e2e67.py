"""
Modèles de rapport et documents générés (plan d'assemblage FE0594) : modèle livré importé comme modèle par défaut,
PDF produit à la consultation (jamais en arrière-plan) puis conservé, envoi d'une nouvelle version (refus motivés,
brouillon, aperçu), modèle conçu pour une version de formulaire envoyé APRÈS la création du document, révision du logo
pour les seuls documents créés après, réimpression avec le modèle actuel, nouvelle version obligatoire quand les
colonnes changent, intégrité en base (versions, affectations, ordre de création).
Nécessite le service de conversion Gotenberg (http://localhost:3000 en développement).
"""
from copy import copy
import base64, hashlib, hmac, io, json, re, time, urllib.request, urllib.error, uuid, zipfile
from pathlib import Path
import openpyxl
from configuration import BASE, SECRET, sql


def b64(d): return base64.urlsafe_b64encode(d).rstrip(b'=').decode()
def jwt(mat, role):
    h = b64(json.dumps({'alg': 'HS256', 'typ': 'JWT'}).encode()); now = int(time.time())
    p = b64(json.dumps({'sub': '00000000-0000-0000-0000-0000000' + mat, 'stamp': '11111111-2222-3333-4444-555555555555', 'matricule': mat,
        'http://schemas.microsoft.com/ws/2008/06/identity/claims/role': role,
        'iss': 'SopalTraceApi', 'aud': 'SopalTraceVueJs', 'iat': now, 'nbf': now, 'exp': now + 3600}).encode())
    return f'{h}.{p}.' + b64(hmac.new(SECRET, f'{h}.{p}'.encode(), hashlib.sha256).digest())
def requete(tok, method, path, corps=None, type_contenu='application/json'):
    req = urllib.request.Request(BASE + path, method=method, data=corps, headers={'Authorization': 'Bearer ' + tok, 'Content-Type': type_contenu})
    try:
        with urllib.request.urlopen(req, timeout=120) as r:
            return r.status, r.headers.get('Content-Type', ''), r.read()
    except urllib.error.HTTPError as e:
        return e.code, e.headers.get('Content-Type', ''), e.read()
def call(tok, method, path, body=None):
    st, _, t = requete(tok, method, path, json.dumps(body).encode() if body is not None else None)
    t = t.decode()
    return st, (json.loads(t) if t.strip().startswith(('{', '[')) else t[:300])
def data(r): return r.get('data', r) if isinstance(r, dict) else r
def envoyer(tok, code, nom, contenu, commentaire=None, formulaire=None):
    limite = 'e2e' + uuid.uuid4().hex
    corps = (f'--{limite}\r\nContent-Disposition: form-data; name="fichier"; filename="{nom}"\r\n'
             f'Content-Type: application/octet-stream\r\n\r\n').encode() + contenu + b'\r\n'
    if commentaire:
        corps += f'--{limite}\r\nContent-Disposition: form-data; name="commentaire"\r\n\r\n{commentaire}\r\n'.encode()
    if formulaire:
        corps += f'--{limite}\r\nContent-Disposition: form-data; name="formulaireId"\r\n\r\n{formulaire}\r\n'.encode()
    corps += f'--{limite}--\r\n'.encode()
    st, _, t = requete(tok, 'POST', f'/rapports/modeles/{code}/versions', corps, f'multipart/form-data; boundary={limite}')
    t = t.decode()
    return st, (json.loads(t) if t.strip().startswith('{') else {})

RQ = jwt('22222', 'RESPONSABLE_QUALITE'); SUP = jwt('44444', 'SUPERVISEUR_QUALITE'); OP = jwt('55555', 'OPERATEUR'); ADMIN = jwt('66666', 'ADMIN')
GEST = ADMIN  # gestion des modèles : l'administrateur (le superviseur qualité définit les formulaires)
ok = ko = 0
def check(label, cond, info=''):
    global ok, ko
    if cond: ok += 1; print('OK ', label)
    else: ko += 1; print('KO ', label, str(info)[:400])
def attendre(condition, secondes=40):
    fin = time.time() + secondes
    while time.time() < fin:
        valeur = condition()
        if valeur: return valeur
        time.sleep(1)
    return condition()
def nb(requete_sql): return int(sql(requete_sql) or 0)

MODELE_LIVRE = (Path(__file__).resolve().parents[1] / 'fixtures' / 'modeles' / 'PLAN_ASS.xlsx').read_bytes()
def variante(modifier):
    classeur = openpyxl.load_workbook(io.BytesIO(MODELE_LIVRE)); feuille = classeur.active
    modifier(feuille)
    sortie = io.BytesIO(); classeur.save(sortie); return sortie.getvalue()

# ── A. Modèles livrés importés : un modèle dédié à chaque formulaire de départ ──────────────
FORMULAIRE = 'FE-ASS-RBGFA-BAC-01'
GROUPE = sql(f"SELECT TOP 1 LOWER(CAST(VersionGroupId AS VARCHAR(36))) FROM Ref_Formulaire WHERE CodeReference='{FORMULAIRE}'")
version1 = attendre(lambda: sql("SELECT Id FROM ReportTemplateVersion WHERE ReportTemplateCode='PLAN_ASS' AND VersionNumber=1 AND Status='PUBLISHED' "
                                f"AND UploadedBy IS NULL AND FormGroupId='{GROUPE}'"), 30)
check("A1 modèle livré PLAN_ASS importé au démarrage : version 1 publiée, dédiée à FE-ASS-RBGFA-BAC-01 (envoyée par l'application)",
      bool(version1) and nb(f"SELECT COUNT(*) FROM ReportTemplateVersion v JOIN Ref_Formulaire f ON f.VersionGroupId = v.FormGroupId AND f.Statut <> 'ARCHIVE' "
                            f"WHERE v.Id='{version1}' AND v.Name = f.Designation + NCHAR(32) + NCHAR(8212) + NCHAR(32) + '{FORMULAIRE}.0.1'") == 1)
check("A2 rattachement de départ : la version 1 est le modèle du formulaire FE-ASS-RBGFA-BAC-01",
      attendre(lambda: sql(f"SELECT a.Kind FROM ReportTemplateAssignment a JOIN Ref_Formulaire f ON f.Id = a.FormId "
                           f"WHERE a.ReportTemplateVersionId='{version1}' AND f.CodeReference='{FORMULAIRE}'") == 'FORM_STRUCTURE', 30))
check('A3 tous les rattachements de départ faits (un par formulaire listé, aucun modèle par défaut pour un modèle de formulaire)',
      attendre(lambda: nb("SELECT COUNT(DISTINCT i.FormCodeReference) FROM ReportTemplateInitialForm i JOIN Ref_Formulaire f ON f.CodeReference = i.FormCodeReference "
         "JOIN ReportTemplateAssignment a ON a.FormId = f.Id") == nb("SELECT COUNT(*) FROM ReportTemplateInitialForm"), 30)
      and nb("SELECT COUNT(*) FROM ReportTemplateAssignment a JOIN ReportTemplateVersion v ON v.Id = a.ReportTemplateVersionId "
             "JOIN ReportTemplate t ON t.Code = v.ReportTemplateCode WHERE a.Kind = 'DEFAULT' AND t.FormRole IS NOT NULL") == 0)
check('A3b modèles dédiés : aucune version partagée entre deux formulaires, chacune nommée pour son formulaire',
      nb("SELECT COUNT(*) FROM (SELECT a.ReportTemplateVersionId FROM ReportTemplateAssignment a JOIN Ref_Formulaire f ON f.Id = a.FormId "
         "GROUP BY a.ReportTemplateVersionId HAVING COUNT(DISTINCT f.VersionGroupId) > 1) x") == 0
      and nb("SELECT COUNT(*) FROM ReportTemplateVersion v JOIN ReportTemplate t ON t.Code = v.ReportTemplateCode "
             "WHERE v.FormGroupId IS NOT NULL AND v.Name NOT LIKE '%' + NCHAR(32) + NCHAR(8212) + NCHAR(32) + '%.' + CAST(v.VersionNumber AS VARCHAR(10))") == 0)
check('A4 empreinte du contenu calculée par la base (SHA-256 du fichier livré)',
      sql(f"SELECT ContentSha256 FROM ReportTemplateVersion WHERE Id='{version1}'") == hashlib.sha256(MODELE_LIVRE).hexdigest().upper())

st_liste, _ = call(SUP, 'GET', '/rapports/modeles')
st_publier, _ = call(SUP, 'POST', f'/rapports/modeles/versions/{version1}/par-defaut', {})
st_suivi, _ = call(SUP, 'GET', '/rapports/formulaires')
check("A5a modèles Excel gérés par l'administrateur : le superviseur qualité suit leur état mais ne les gère pas",
      (st_liste, st_publier, st_suivi) == (403, 403, 200), (st_liste, st_publier, st_suivi))
st, r = call(GEST, 'GET', '/rapports/modeles')
modeles = {m['code']: m for m in (data(r) or [])}
check('A5 liste des modèles : nom avec le code Sopal, rôle des formulaires, formulaires rattachés',
      st == 200 and modeles['PLAN_ASS']['nom'].endswith('(FE0594)') and modeles['PLAN_ASS']['formRole'] == 'EN_COURS_DE_ASSEMBLAGE'
      and f'{FORMULAIRE} version 0' in modeles['PLAN_ASS']['formulaires']
      and 'FE-VM-MAS22 version 0' in modeles['VERIF_MACHINE_MAS22']['formulaires']
      and 'FE-RC-ENCF-ASS-AUTO version 0' in modeles['PLAN_ASS_RESULTATS']['formulaires']
      and 'FE-RC-ENCF-ASS-PAS78_SOUPAPE version 0' in modeles['PLAN_ASS_RESULTATS_PAS72']['formulaires']
      and modeles['REGISTRE_TRACA']['formRole'] == 'REGISTRE_TRACABILITE' and 'FE0129 version 0' in modeles['REGISTRE_TRACA']['formulaires'], (st, r))
check('A6 machine sans modèle rattaché (MAS20) : signalée parmi les formulaires sans modèle',
      'FE-VM-MAS20 version 0' in modeles['VERIF_MACHINE']['formulairesSansModele'])
check('A7 opérateur, responsable qualité et superviseur qualité : gestion des modèles refusée (403)',
      all(call(t, 'GET', '/rapports/modeles')[0] == 403 for t in (OP, RQ, SUP)))
st, r = call(GEST, 'GET', '/rapports/modeles/PLAN_ASS/champs')
noms = {c['nom']: c for c in data(r) or []}
check('A8 champs disponibles : contrat du plan (document, section, ligne) avec la syntaxe à écrire',
      st == 200 and noms.get('Ligne.Caracteristique', {}).get('syntaxe') == '{{Ligne.Caracteristique}}' and noms.get('Section.Libelle', {}).get('portee') == 'SECTION', (st, r))

# ── B. Plan d'assemblage : PDF produit à la consultation ─────────────────────
famille = sql("SELECT TOP 1 Code FROM FamilleProduitFini WHERE Code='RBGFA-BAC-01'")
def ligne(ordre, libelle, limite, moyen, frequence, test=None, observations=None):
    extra = [{'cleColonne': 'FREQUENCE', 'valeurColonne': frequence}] + ([{'cleColonne': 'TEST', 'valeurColonne': test}] if test else [])
    return {'ordreAffiche': ordre, 'libelleAffiche': libelle, 'limiteSpecTexte': limite, 'moyenTexteLibre': moyen,
            'observations': observations, 'extraColonnes': extra}
def sections_plan(avec_test=False):
    t = (lambda v: v) if avec_test else (lambda v: None)
    return [
        {'ordreAffiche': 1, 'libelleSection': 'Caractéristiques à contrôler aux réglages (une série de 12 pièces)',
         'lignes': [ligne(1, 'Étanchéité du clapet', '0 fuite', "Banc d'étanchéité", 'Chaque série', t('T1'), 'Arrêt si fuite'),
                    ligne(2, 'Couple de serrage', '12 ± 1 N.m', 'Clé dynamométrique', '5 pièces / heure', t('T2'))]},
        {'ordreAffiche': 2, 'libelleSection': 'Caractéristiques à contrôler en cours de production',
         'lignes': [ligne(1, 'Aspect visuel', 'Sans rayure', 'Visuel', '100 %', t('T3'))]}]
COLONNES = [{'key': 'FREQUENCE', 'label': 'Fréquence', 'type': 'TEXT'}]
COLONNES_TEST = COLONNES + [{'key': 'TEST', 'label': 'Test', 'type': 'TEXT'}]
st, r = call(SUP, 'POST', '/Document', {
    'typeDocumentCode': 'PLAN_ASS', 'nom': 'Plan RBGFA-BAC-01 E2E rapport', 'familleProduitFiniCode': famille,
    'refFormulaireCodeReference': FORMULAIRE, 'remarques': 'NB : Pour chaque lot, vérifier les fiches suiveuses.',
    'colonneDefs': COLONNES, 'sections': sections_plan()})
doc = (data(r) or {}).get('id') if isinstance(data(r), dict) else None
check("B1 plan d'assemblage créé par le superviseur sur FE-ASS-RBGFA-BAC-01 (colonne propre FREQUENCE)", st in (200, 201) and doc, (st, r))
formulaire = sql(f"SELECT FormulaireId FROM Document_Entete WHERE Id='{doc}'").lower()
check("B2 ordre de création attribué par la base", nb(f"SELECT CreationOrder FROM Document_Entete WHERE Id='{doc}'") > 0)
time.sleep(3)
check('B3 aucun PDF produit en arrière-plan', nb(f"SELECT COUNT(*) FROM GeneratedDocument WHERE DocumentEnteteId='{doc}'") == 0)

etat = lambda d=None: data(call(SUP, 'GET', f'/rapports/documents/par-document/{d or doc}')[1]) or {}
e = etat()
check('B4 état : modèle du formulaire (version 1), en vigueur à la création, aucun PDF encore',
      (e.get('source', {}).get('modele') or {}).get('numeroVersion') == 1
      and e['source']['modele']['regle'] == 'EN_VIGUEUR_A_LA_CREATION' and e.get('documents') == [] and e.get('modeleActuel') is None, e)

st, type_pdf, pdf = requete(SUP, 'GET', f'/rapports/documents/par-document/{doc}/pdf')
genere = sql(f"SELECT Id FROM GeneratedDocument WHERE DocumentEnteteId='{doc}'")
check('B5 consultation : PDF produit au clic (signature %PDF), avec la version 1, demandeur tracé',
      st == 200 and type_pdf.startswith('application/pdf') and pdf[:5] == b'%PDF-'
      and sql(f"SELECT CONCAT(ReportTemplateVersionId, '|', RequestedBy) FROM GeneratedDocument WHERE Id='{genere}'") == f'{version1}|44444', (st, pdf[:20]))
check('B6 empreinte du PDF calculée par la base = SHA-256 du fichier reçu',
      sql(f"SELECT PdfSha256 FROM GeneratedDocument WHERE Id='{genere}'") == hashlib.sha256(pdf).hexdigest().upper())
st, _, pdf_rq = requete(RQ, 'GET', f'/rapports/documents/par-document/{doc}/pdf')
check('B7 deuxième consultation (responsable qualité) : même PDF, rien de reproduit',
      st == 200 and pdf_rq == pdf and nb(f"SELECT COUNT(*) FROM GeneratedDocument WHERE DocumentEnteteId='{doc}'") == 1)
st, _, xlsx = requete(SUP, 'GET', f'/rapports/documents/par-document/{doc}/excel')
check('B8 Excel rempli téléchargeable (classeur .xlsx)', st == 200 and zipfile.is_zipfile(io.BytesIO(xlsx)))
feuille = openpyxl.load_workbook(io.BytesIO(xlsx)).active
valeurs = [[c.value for c in rangee] for rangee in feuille.iter_rows(min_row=1, max_row=15, max_col=6)]
check('B9 Excel : 2 sections, leurs lignes, remarques sous le tableau, aucun champ {{}} restant',
      feuille['A6'].value.startswith('Caractéristiques à contrôler aux réglages') and feuille['A8'].value == 'Étanchéité du clapet'
      and feuille['B9'].value == '12 ± 1 N.m' and feuille['A10'].value == 'Caractéristiques à contrôler en cours de production'
      and feuille['D12'].value == 'Visuel' and feuille['A13'].value.startswith('NB : Pour chaque lot')
      and not any('{{' in str(v) for rangee in valeurs for v in rangee if v), valeurs)
check('B10 Excel : date, code article « type de robinet : famille » et pied de page « code formulaire.version » remplis',
      feuille['A3'].value.startswith('Date : ') and feuille['C3'].value == f'Code Article : AUTO : {famille}'
      and f'{FORMULAIRE}.0' in (feuille.oddFooter.right.text or ''), (feuille['A3'].value, feuille['C3'].value, feuille.oddFooter.right.text))
def taille_logo(classeur):
    archive = zipfile.ZipFile(io.BytesIO(classeur))
    dessin = next(n for n in archive.namelist() if n.startswith('xl/drawings/drawing') and n.endswith('.xml'))
    return re.search(r'ext cx="(\d+)" cy="(\d+)"', archive.read(dessin).decode()).groups()
check('B11 logo à la taille du modèle (non agrandi)', taille_logo(xlsx) == taille_logo(MODELE_LIVRE), (taille_logo(xlsx), taille_logo(MODELE_LIVRE)))
st, _, modele = requete(SUP, 'GET', f'/rapports/documents/par-document/{doc}/modele')
check('B12 superviseur : Excel du modèle du document = fichier de la version 1', st == 200 and modele == MODELE_LIVRE)
check('B13 opérateur : PDF refusé (403) ; responsable qualité : Excel du modèle refusé (403)',
      requete(OP, 'GET', f'/rapports/documents/par-document/{doc}/pdf')[0] == 403
      and requete(RQ, 'GET', f'/rapports/documents/par-document/{doc}/modele')[0] == 403)

# ── C. Envoi d'une nouvelle version : refus motivés ──────────────────────────
versions_avant = nb("SELECT COUNT(*) FROM ReportTemplateVersion WHERE ReportTemplateCode='PLAN_ASS'")
st, r = envoyer(GEST, 'PLAN_ASS', 'sans_formulaire.xlsx', MODELE_LIVRE)
check('C0 modèle de formulaire envoyé sans formulaire visé : refusé (un modèle est dédié à un formulaire)',
      st == 400 and any('Choisissez le formulaire' in e for e in r.get('erreurs', [])), (st, r))
st, r = envoyer(GEST, 'PLAN_ASS', 'faute.xlsx', variante(lambda f: f.__setitem__('A8', '{{Ligne.Caracteristiqe}}')), formulaire=formulaire)
check('C1 champ mal écrit : refusé avec la cellule et le nom voisin proposé',
      st == 400 and any('Cellule A8' in e and '{{Ligne.Caracteristique}}' in e for e in r.get('erreurs', [])), (st, r))
st, r = envoyer(GEST, 'PLAN_ASS', 'colonne.xlsx', variante(lambda f: f.__setitem__('G8', '{{Ligne.Colonne.FREQENCE}}')), formulaire=formulaire)
check('C2 colonne propre inconnue des formulaires : refusée', st == 400 and any('FREQENCE' in e for e in r.get('erreurs', [])), (st, r))
st, r = envoyer(GEST, 'PLAN_ASS', 'ligne.xlsx', variante(lambda f: f.__setitem__('G8', '{{Ligne.Frequence}}')), formulaire=formulaire)
check('C3 colonne propre écrite comme un champ de ligne ({{Ligne.Frequence}}) : refusée avec la bonne écriture',
      st == 400 and any('{{Ligne.Colonne.FREQUENCE}}' in e for e in r.get('erreurs', [])), (st, r))
st, r = envoyer(GEST, 'PLAN_ASS', 'deux.xlsx', variante(lambda f: f.__setitem__('A5', '{{Ligne.Observations}}')), formulaire=formulaire)
check('C4 champs Ligne.* sur deux lignes Excel : refusé (structure du bloc répété)', st == 400 and any('une seule ligne' in e for e in r.get('erreurs', [])), (st, r))
st, r = envoyer(GEST, 'PLAN_ASS', 'modele.pdf', b'%PDF-1.4 pas un classeur', formulaire=formulaire)
check("C5 fichier qui n'est pas un .xlsx : refusé", st == 400 and r.get('erreurs'), (st, r))
nb_versions_plan_ass = nb("SELECT COUNT(*) FROM ReportTemplateVersion WHERE ReportTemplateCode='PLAN_ASS'")
st, r = envoyer(GEST, 'PLAN_ASS', 'PLAN_ASS.xlsx', MODELE_LIVRE, formulaire=formulaire)
check('C6 fichier identique à une version du modèle de ce formulaire : version 1 reprise, aucune nouvelle version',
      st == 200 and ((r.get('data') or {}).get('version') or {}).get('numero') == 1
      and any('déjà la version 1' in e for e in (r.get('data') or {}).get('avertissements', []))
      and nb("SELECT COUNT(*) FROM ReportTemplateVersion WHERE ReportTemplateCode='PLAN_ASS'") == nb_versions_plan_ass, (st, r))
EXCEL_VIDE = (Path(__file__).resolve().parents[1] / 'fixtures' / 'excel' / 'Plan_de_controle_FE0594.xlsx').read_bytes()
st, r = envoyer(GEST, 'PLAN_ASS', 'Plan_de_controle_FE0594.xlsx', EXCEL_VIDE, formulaire=formulaire)
check("C7 Excel vide d'origine (sans les champs du tableau) : refusé, champs obligatoires signalés",
      st == 400 and any('{{Ligne.Caracteristique}}' in e for e in r.get('erreurs', [])), (st, r))
check("C8 aucun fichier refusé n'est enregistré", nb("SELECT COUNT(*) FROM ReportTemplateVersion WHERE ReportTemplateCode='PLAN_ASS'") == versions_avant)
check("C9 responsable qualité et superviseur qualité : envoi d'un modèle refusé (403)",
      all(envoyer(t, 'PLAN_ASS', 'x.xlsx', MODELE_LIVRE)[0] == 403 for t in (RQ, SUP)))

# ── D. Colonne ajoutée : nouvelle version du formulaire SANS modèle, puis modèle rattaché ──
def nouvelle_version(ancien, colonnes, sections, remarques):
    courant = data(call(SUP, 'GET', f'/Document/{ancien}')[1]) or {}
    st, r = call(SUP, 'POST', f'/Document/{ancien}/version', {
        'ancienId': ancien, 'rowVersion': courant.get('rowVersion'), 'typeDocumentCode': 'PLAN_ASS', 'nom': 'Plan RBGFA-BAC-01 E2E rapport',
        'familleProduitFiniCode': famille, 'refFormulaireCodeReference': FORMULAIRE, 'remarques': remarques,
        'colonneDefs': colonnes, 'sections': sections})
    return st, r, ((data(r) or {}).get('id') if isinstance(data(r), dict) else None)
st, r, doc2 = nouvelle_version(doc, COLONNES_TEST, sections_plan(avec_test=True), 'NB : colonne Test ajoutée.')
formulaire2 = sql(f"SELECT FormulaireId FROM Document_Entete WHERE Id='{doc2}'").lower() if doc2 else ''
check('D1 nouvelle version du plan avec la colonne Test : nouvelle version du formulaire (version 1)',
      st in (200, 201) and doc2 and sql(f"SELECT Version FROM Ref_Formulaire WHERE Id='{formulaire2}'") == '1', (st, r))
check("D1b colonnes changées : demande automatique « nouvelle structure » envoyée à l'administrateur",
      sql(f"SELECT CONCAT(Origin, '|', Status) FROM ReportTemplateRequest WHERE FormId='{formulaire2}'") == 'NEW_STRUCTURE|PENDING')
e2 = etat(doc2)
check("D2 colonnes changées : aucun modèle hérité de la version 0, modèle à rattacher (message clair)",
      e2.get('source', {}).get('modele') is None and e2['source'].get('modeleManquant') == 'STRUCTURE_MODIFIEE'
      and 'a changé de colonnes' in (e2['source'].get('messageModeleManquant') or ''), e2)
st, _, corps = requete(SUP, 'GET', f'/rapports/documents/par-document/{doc2}/pdf')
check('D3 PDF refusé tant que le modèle n\'est pas rattaché (le document, lui, est enregistré)',
      st == 400 and 'a changé de colonnes' in json.loads(corps).get('message', ''), (st, corps[:300]))
st, r = call(GEST, 'GET', '/rapports/modeles')
plan = next(m for m in data(r) if m['code'] == 'PLAN_ASS')
check('D4 écran des modèles : la version 1 du formulaire est signalée sans modèle', f'{FORMULAIRE} version 1' in plan['formulairesSansModele'], plan)

def v2(f):
    f['G7'] = 'Fréquence'; f['G8'] = '{{Ligne.Colonne.FREQUENCE}}'
    f['H7'] = 'Test'; f['H8'] = '{{Ligne.Colonne.TEST}}'
    for c in ('G7', 'H7'): f[c].font = copy(f['F7'].font); f[c].border = copy(f['F7'].border)
    for c in ('G8', 'H8'): f[c].border = copy(f['F8'].border)
    f.print_area = 'A1:H9'
st, r = envoyer(GEST, 'PLAN_ASS', 'Plan_de_controle_FE0594_v2.xlsx', variante(v2), 'Colonnes Fréquence et Test', formulaire)
check("D5 modèle visant la version 0 du formulaire (archivée par la nouvelle version) : refusé, ses documents gardent leur modèle",
      st == 400 and 'archivé' in json.dumps(r, ensure_ascii=False), (st, r))
st, r = envoyer(GEST, 'PLAN_ASS', 'Plan_de_controle_FE0594_v2.xlsx', variante(v2), 'Colonnes Fréquence et Test', formulaire2)
version2 = ((data(r) or {}).get('version') or {}).get('id')
check('D6 modèle vérifié contre la version 1 du formulaire (qui a la colonne Test) : brouillon numéro 2',
      st == 200 and version2 and sql(f"SELECT CONCAT(VersionNumber,'/',Status) FROM ReportTemplateVersion WHERE Id='{version2}'") == '2/DRAFT', (st, r))
st, type_pdf, apercu = requete(GEST, 'GET', f'/rapports/modeles/versions/{version2}/apercu?formulaireId={formulaire2}')
check("D7 aperçu PDF du brouillon (modèle vierge avec l'identité du formulaire), jamais archivé",
      st == 200 and apercu[:5] == b'%PDF-' and nb(f"SELECT COUNT(*) FROM GeneratedDocument WHERE ReportTemplateVersionId='{version2}'") == 0, (st, apercu[:200]))
st, r = call(GEST, 'POST', f'/rapports/modeles/versions/{version2}/formulaire', {'formulaireId': formulaire2, 'commentaire': 'Structure avec Test'})
check('D8b la demande « nouvelle structure » est traitée par ce rattachement (rattachement gardé dans la demande)',
      sql(f"SELECT CONCAT(r.Status, '|', v.Id) FROM ReportTemplateRequest r JOIN ReportTemplateAssignment a ON a.Id = r.AssignmentId "
          f"JOIN ReportTemplateVersion v ON v.Id = a.ReportTemplateVersionId WHERE r.FormId='{formulaire2}'").lower() == f'done|{version2}'.lower())
check('D8 version 2 rattachée à la version 1 du formulaire : publiée, rattachement tracé',
      st == 200 and sql(f"SELECT CONCAT(Status,'/',PublishedBy) FROM ReportTemplateVersion WHERE Id='{version2}'") == 'PUBLISHED/66666'
      and sql(f"SELECT CONCAT(Kind,'/',AssignedBy) FROM ReportTemplateAssignment WHERE ReportTemplateVersionId='{version2}'") == 'FORM_STRUCTURE/66666', (st, r))
e2 = etat(doc2)
check('D9 le plan créé AVANT ce modèle le reçoit (modèle de sa version de formulaire)',
      (e2['source'].get('modele') or {}).get('numeroVersion') == 2 and e2['source']['modele']['regle'] == 'CONCU_POUR_LA_VERSION_DU_FORMULAIRE', e2)
_, _, xlsx2 = requete(SUP, 'GET', f'/rapports/documents/par-document/{doc2}/excel')
f2 = openpyxl.load_workbook(io.BytesIO(xlsx2)).active
check('D10 nouvelles colonnes remplies depuis les colonnes propres du plan, sans développement',
      (f2['G8'].value, f2['H8'].value, f2['H9'].value, f2['H12'].value) == ('Chaque série', 'T1', 'T2', 'T3'), [f2[f'H{i}'].value for i in range(6, 14)])
check("D11 le plan précédent (formulaire version 0) garde la version 1 du modèle", (etat()['source'].get('modele') or {}).get('numeroVersion') == 1)
st, r = call(GEST, 'POST', f'/rapports/modeles/versions/{version2}/formulaire', {'formulaireId': formulaire2})
check('D12 même version déjà rattachée à ce formulaire : refusé', st == 400, (st, r))

# ── E. Mêmes colonnes : la nouvelle version garde le modèle ; un autre peut lui être rattaché ──
st, r, doc3 = nouvelle_version(doc2, COLONNES_TEST, sections_plan(avec_test=True), 'NB : contenu des lignes corrigé.')
formulaire3 = sql(f"SELECT FormulaireId FROM Document_Entete WHERE Id='{doc3}'").lower() if doc3 else ''
e3 = etat(doc3)
check('E1 nouvelle version aux mêmes colonnes : le modèle de la version précédente (M2) est recopié explicitement (« repris de »), sans demande',
      st in (200, 201) and (e3.get('source', {}).get('modele') or {}).get('numeroVersion') == 2
      and e3['source']['modele']['regle'] == 'EN_VIGUEUR_A_LA_CREATION'
      and sql(f"SELECT CONCAT(Kind, '|', ReportTemplateVersionId) FROM ReportTemplateAssignment WHERE FormId='{formulaire3}'").lower() == f'form_structure|{version2}'.lower()
      and 'Repris de FE-ASS-RBGFA-BAC-01.1' in sql(f"SELECT Comment FROM ReportTemplateAssignment WHERE FormId='{formulaire3}'")
      and nb(f"SELECT COUNT(*) FROM ReportTemplateRequest WHERE FormId='{formulaire3}'") == 0, e3)
def v3(f):
    v2(f); f.oddHeader.left.text = 'Révision du logo'
st, r = envoyer(GEST, 'PLAN_ASS', 'Plan_de_controle_FE0594_v3.xlsx', variante(v3), 'Nouveau logo', formulaire3)
version3 = ((data(r) or {}).get('version') or {}).get('id')
st, r = call(GEST, 'POST', f'/rapports/modeles/versions/{version3}/revision', {'versionRemplaceeId': version2, 'commentaire': 'Nouveau logo'})
check("E2 révision (logo) de la version 2 : publiée pour chaque formulaire dont les prochains documents l'utilisaient",
      st == 200 and sql(f"SELECT COUNT(*) FROM ReportTemplateAssignment WHERE ReportTemplateVersionId='{version3}' AND Kind='FORM_REVISION'") == '1', (st, r))
e3 = etat(doc3)
check('E3 le plan déjà créé garde son modèle (version 2) ; modèle actuel proposé : version 3',
      (e3['source'].get('modele') or {}).get('numeroVersion') == 2 and (e3.get('modeleActuel') or {}).get('numeroVersion') == 3, e3)
st, _, reimpression = requete(SUP, 'GET', f'/rapports/documents/par-document/{doc3}/pdf?modeleActuel=true')
check("E4 réimpression avec le modèle actuel à la demande : PDF avec la version 3, l'original reste",
      st == 200 and reimpression[:5] == b'%PDF-' and nb(f"SELECT COUNT(*) FROM GeneratedDocument g WHERE g.DocumentEnteteId='{doc3}' "
                                                         f"AND g.ReportTemplateVersionId='{version3}' AND g.FileName LIKE '%reimpression%'") == 1)
st, r, doc4 = nouvelle_version(doc3, COLONNES_TEST, sections_plan(avec_test=True), 'NB : version suivante.')
check('E5 le plan créé après la révision prend la version 3', (etat(doc4).get('source', {}).get('modele') or {}).get('numeroVersion') == 3)

# ── F. Structure modifiée en correction mineure : refusée ────────────────────
courant = data(call(SUP, 'GET', f'/Document/{doc4}')[1]) or {}
modif = {'rowVersion': courant.get('rowVersion'), 'nom': courant.get('nom'), 'remarques': 'NB : remarque corrigée.', 'sections': courant.get('sections')}
st, r = call(SUP, 'PUT', f'/Document/{doc4}', dict(modif, configurationColonnesJson=json.dumps(COLONNES_TEST + [{'key': 'AUTRE', 'label': 'Autre', 'type': 'TEXT'}])))
check('F1 correction mineure qui ajoute une colonne : refusée (nouvelle version obligatoire)',
      st == 400 and 'nouvelle version' in json.dumps(r, ensure_ascii=False), (st, r))
check("F2 la colonne refusée n'a pas été ajoutée au formulaire",
      nb(f"SELECT COUNT(*) FROM Ref_Formulaire_ColonneDef c JOIN Document_Entete d ON d.FormulaireId = c.FormulaireId "
         f"WHERE d.Id='{doc4}' AND c.CleColonne='AUTRE'") == 0)
st, r = call(SUP, 'PUT', f'/Document/{doc4}', dict(modif, configurationColonnesJson=json.dumps(COLONNES_TEST)))
check('F3 correction mineure sans changement de colonnes : acceptée, aucun PDF produit', st == 200
      and nb(f"SELECT COUNT(*) FROM GeneratedDocument WHERE DocumentEnteteId='{doc4}'") == 0, (st, r))

# ── G. Modèle par défaut réservé aux documents sans formulaire ───────────────
LOGO_CELLULE = (Path(__file__).resolve().parents[2] / 'SopalTrace.Application.Tests' / 'Rapports' / 'Fixtures' / 'modele_logo_dans_cellule.xlsx').read_bytes()
formulaire4 = sql(f"SELECT LOWER(CAST(FormulaireId AS VARCHAR(36))) FROM Document_Entete WHERE Id='{doc4}'")
st, r = envoyer(GEST, 'PLAN_ASS', 'PLAN_ASS_logo_cellule.xlsx', LOGO_CELLULE, 'Logo dans la cellule', formulaire4)
version_logo = ((data(r) or {}).get('version') or {}).get('id')
check('G1 modèle avec logo « dans la cellule » (Excel 365) accepté', st == 200 and version_logo, (st, r))
st, _, apercu = requete(GEST, 'GET', f'/rapports/modeles/versions/{version_logo}/apercu')
check('G2 aperçu PDF du modèle avec logo dans la cellule', st == 200 and apercu[:5] == b'%PDF-', (st, apercu[:200]))
st, r = call(GEST, 'POST', f'/rapports/modeles/versions/{version_logo}/par-defaut', {})
check('G3 un modèle de plan ne devient pas « par défaut » : il se rattache à chaque formulaire', st == 400 and 'rattachez' in json.dumps(r, ensure_ascii=False), (st, r))
vm = sql("SELECT TOP 1 Id FROM Ref_Formulaire WHERE CodeReference='FE-VM-MAS20'").lower()
st, r = call(GEST, 'POST', f'/rapports/modeles/versions/{version2}/formulaire', {'formulaireId': vm})
check("G4 modèle de plan rattaché au formulaire d'une machine : refusé (autre rôle)", st == 400 and 'pas un formulaire' in json.dumps(r, ensure_ascii=False), (st, r))

# ── H. Liste des documents (superviseur) ─────────────────────────────────────
st, r = call(SUP, 'GET', '/rapports/documents?categorie=QUALITE&type=PLAN_ASS')
liste = {d.get('documentEnteteId'): d for d in data(r) or []}
check('H1 liste des documents qualité : chaque plan avec son modèle', st == 200 and all(liste.get(d, {}).get('modele') for d in (doc, doc2, doc3, doc4))
      and liste[doc2]['formulaire'] == f'{FORMULAIRE} version 1', (st, list(liste.values())[:2]))
st, r = call(SUP, 'GET', '/rapports/documents?categorie=QUALITE')
sans_modele = next((d for d in data(r) or [] if not d.get('contratDisponible')), None)
if sans_modele:
    st, _, corps = requete(SUP, 'GET', f"/rapports/documents/par-document/{sans_modele['documentEnteteId']}/pdf")
    check("H2 document dont l'Excel n'est pas encore fourni : message clair", st == 400 and "Excel" in json.loads(corps).get('message', ''), (st, corps[:300]))

# ── I. Intégrité garantie par la base ────────────────────────────────────────
def refuse(requete_sql):
    sortie = sql(requete_sql)
    return 'Msg' in sortie or 'interdit' in sortie or 'ne peut' in sortie
check('I1 document généré : modification refusée par la base', refuse(f"UPDATE GeneratedDocument SET FileName='x' WHERE Id='{genere}'"))
check('I2 document généré : suppression refusée par la base', refuse(f"DELETE FROM GeneratedDocument WHERE Id='{genere}'")
      and nb(f"SELECT COUNT(*) FROM GeneratedDocument WHERE Id='{genere}'") == 1)
check("I3 fichier d'une version de modèle : modification refusée", refuse(f"UPDATE ReportTemplateVersion SET Content=0x00 WHERE Id='{version1}'"))
check('I4 version de modèle : suppression refusée', refuse(f"DELETE FROM ReportTemplateVersion WHERE Id='{version1}'"))
check('I5 version publiée : retour en brouillon refusé', refuse(f"UPDATE ReportTemplateVersion SET Status='DRAFT', PublishedAt=NULL WHERE Id='{version1}'"))
check('I6 historique des rattachements : modification et suppression refusées',
      refuse(f"UPDATE ReportTemplateAssignment SET ReportTemplateVersionId='{version2}' WHERE ReportTemplateVersionId='{version1}'")
      and refuse(f"DELETE FROM ReportTemplateAssignment WHERE ReportTemplateVersionId='{version1}'"))
check("I7 rattachement d'un brouillon : refusé", refuse(
    f"INSERT INTO ReportTemplateAssignment (ReportTemplateVersionId, Kind, FormId) VALUES ('{version_logo}', 'FORM_STRUCTURE', '{formulaire3}')"))
check("I8 rattachement à une version de formulaire archivée : refusé",
      sql(f"SELECT Statut FROM Ref_Formulaire WHERE Id='{formulaire}'") == 'ARCHIVE'
      and refuse(f"INSERT INTO ReportTemplateAssignment (ReportTemplateVersionId, FormId, Kind) VALUES ('{version3}', '{formulaire}', 'FORM_REVISION')"))
check("I9 rattachement à un formulaire d'un autre rôle, ou « par défaut » d'un modèle de formulaire : refusés par la base",
      refuse(f"INSERT INTO ReportTemplateAssignment (ReportTemplateVersionId, FormId, Kind) VALUES ('{version3}', '{vm}', 'FORM_STRUCTURE')")
      and refuse(f"INSERT INTO ReportTemplateAssignment (ReportTemplateVersionId, Kind) VALUES ('{version3}', 'DEFAULT')"))
st, r = call(GEST, 'POST', f'/rapports/modeles/versions/{version3}/formulaire', {'formulaireId': formulaire})
check("I10 API : rattachement à une version de formulaire archivée refusé avec un message", st == 400 and 'archivé' in json.dumps(r, ensure_ascii=False), (st, r))
check("I11 ordre de création d'un document : modification refusée", refuse(f"UPDATE Document_Entete SET CreationOrder = CreationOrder + 100000 WHERE Id='{doc}'"))

# ── F. Fiches vierges : la fiche que rempliront les opérateurs, vue depuis le document qualité ──
def contenu_classeur(xlsx):
    classeur = openpyxl.load_workbook(io.BytesIO(xlsx))
    cellules = ' '.join(str(c.value) for f in classeur.worksheets for rangee in f.iter_rows() for c in rangee if c.value is not None)
    pieds = ' '.join(p.text or '' for f in classeur.worksheets for p in (f.oddFooter.left, f.oddFooter.center, f.oddFooter.right))
    return cellules, pieds

def fiche_de(document_id):
    return (data(call(SUP, 'GET', f'/rapports/documents/par-document/{document_id}')[1]) or {}).get('ficheVierge') or {}

mas22 = sql("SELECT LOWER(CAST(Id AS VARCHAR(36))) FROM Document_Entete WHERE TypeDocumentCode='VERIF_MACHINE' AND MachineCode='MAS22' AND Statut='ACTIF'")
fiche = fiche_de(mas22)
check('F1 vérification machine : fiche vierge mise en forme avec le modèle de sa machine (MAS22)',
      (fiche.get('modele') or {}).get('codeModele') == 'VERIF_MACHINE_MAS22' and fiche.get('typeExecution') == 'VERIF_MACHINE', fiche)
st, _, xlsx = requete(SUP, 'GET', f'/rapports/documents/par-document/{mas22}/fiche-vierge/excel')
cellules, pieds = contenu_classeur(xlsx) if st == 200 else ('', '')
check('F2 fiche vierge Excel : machine du plan, aucun champ restant, pied de page code.version du formulaire',
      st == 200 and 'MAS22' in cellules and '{{' not in cellules and 'FE-VM-MAS22.0' in pieds, (st, pieds, cellules[:300]))
st, _, pdf = requete(SUP, 'GET', f'/rapports/documents/par-document/{mas22}/fiche-vierge/pdf')
check('F3 fiche vierge PDF : produite à la demande, jamais conservée dans les documents générés',
      st == 200 and pdf[:4] == b'%PDF' and nb(f"SELECT COUNT(*) FROM GeneratedDocument WHERE DocumentEnteteId='{mas22}'") == 0, st)

fiche = fiche_de(doc)
st, _, xlsx = requete(SUP, 'GET', f'/rapports/documents/par-document/{doc}/fiche-vierge/excel')
cellules, pieds = contenu_classeur(xlsx) if st == 200 else ('', '')
check("F4 plan d'assemblage : fiche des résultats, formulaire de résultats de sa famille, champs remplis sauf les réponses",
      fiche.get('typeExecution') == 'PLAN_ASS' and (fiche.get('modele') or {}).get('codeModele', '').startswith('PLAN_ASS_RESULTATS')
      and st == 200 and '{{' not in cellules and (fiche.get('formulaire') or '').split(' ')[0] in pieds, (fiche, st, pieds))

st, r = call(SUP, 'POST', '/Document', {'typeDocumentCode': 'CTRL_POSTE', 'nom': 'Contrôle poste PAS72 (fiche vierge)', 'posteCode': 'PAS72', 'sections': []})
controle = str((data(r) or {}).get('id', '')).lower() if st in (200, 201) else ''
fiche = fiche_de(controle) if controle else {}
st, _, xlsx = requete(SUP, 'GET', f'/rapports/documents/par-document/{controle}/fiche-vierge/excel')
cellules, pieds = contenu_classeur(xlsx) if st == 200 else ('', '')
check('F5 contrôle au poste : fiche vierge (modèle FE0794), poste du document, pied de page FE-RC-PAS72.version',
      (fiche.get('modele') or {}).get('codeModele') == 'RESULTAT_CONTROLE_POSTE' and st == 200 and 'PAS72' in cellules
      and 'FE-RC-PAS72.' in pieds and '{{' not in cellules, (fiche, st, pieds))

# ── G. Nouveau logo : nouvelle version du formulaire, documents actifs archivés tels quels puis recréés en version suivante ──
actif = sql("SELECT LOWER(CAST(d.Id AS VARCHAR(36))) FROM Document_Entete d JOIN Ref_Formulaire f ON f.Id = d.FormulaireId "
            "WHERE f.CodeReference = 'FE-ASS-RBGFA-BAC-01' AND d.Statut = 'ACTIF' AND d.Portee = 'GENERIQUE'")
ligne_actuelle = sql(f"SELECT LOWER(CAST(FormulaireId AS VARCHAR(36))) FROM Document_Entete WHERE Id='{actif}'")
version_formulaire, version_document, nom_document = sql(
    f"SELECT CONCAT(f.Version, '|', d.Version, '|', d.Nom) FROM Document_Entete d JOIN Ref_Formulaire f ON f.Id = d.FormulaireId WHERE d.Id='{actif}'").split('|')
lignes_document = nb(f"SELECT COUNT(*) FROM Document_Ligne WHERE EnteteId='{actif}'")
modele_avant = (data(call(SUP, 'GET', f'/rapports/documents/par-document/{actif}')[1]) or {}).get('source', {}).get('modele') or {}
st, _, fichier = requete(GEST, 'GET', f"/rapports/modeles/versions/{modele_avant.get('versionId')}/fichier")
classeur = openpyxl.load_workbook(io.BytesIO(fichier)); classeur.active['Z1'] = 'Logo 2026'
sortie = io.BytesIO(); classeur.save(sortie)
st, r = envoyer(GEST, 'PLAN_ASS', 'PLAN_ASS_logo_2026.xlsx', sortie.getvalue(), 'Logo 2026', ligne_actuelle)
version_logo = ((data(r) or {}).get('version') or {}).get('id')
check('G1 nouveau modèle (logo 2026) envoyé en brouillon', st == 200 and version_logo, (st, r))

st, r = call(GEST, 'POST', f'/rapports/modeles/versions/{version_logo}/changement-logo', {'formulaireIds': [ligne_actuelle], 'commentaire': 'Logo 2026'})
nouvelle_ligne = sql(f"SELECT LOWER(CAST(Id AS VARCHAR(36))) FROM Ref_Formulaire WHERE CodeReference = 'FE-ASS-RBGFA-BAC-01' AND Statut <> 'ARCHIVE'")
check('G2 nouveau logo : FE-ASS-RBGFA-BAC-01 passe à la version suivante (même lignée), l\'ancienne version est archivée',
      st == 200 and nouvelle_ligne and nouvelle_ligne != ligne_actuelle
      and sql(f"SELECT CONCAT(Version, ' ', Statut) FROM Ref_Formulaire WHERE Id='{nouvelle_ligne}'") == f'{int(version_formulaire) + 1} ACTIF'
      and sql(f"SELECT Statut FROM Ref_Formulaire WHERE Id='{ligne_actuelle}'") == 'ARCHIVE'
      and nb(f"SELECT COUNT(DISTINCT VersionGroupId) FROM Ref_Formulaire WHERE Id IN ('{ligne_actuelle}', '{nouvelle_ligne}')") == 1, (st, r))
recree = sql(f"SELECT LOWER(CAST(Id AS VARCHAR(36))) FROM Document_Entete WHERE FormulaireId='{nouvelle_ligne}' AND Statut='ACTIF'")
attendu = re.sub(r'V\d+$', f'V{int(version_formulaire) + 1}', nom_document)
check('G3 document archivé tel quel, recréé en version suivante : même contenu, nom et version suivants, motif',
      sql(f"SELECT Statut FROM Document_Entete WHERE Id='{actif}'") == 'ARCHIVE' and bool(recree)
      and sql(f"SELECT CONCAT(Version, '|', Nom) FROM Document_Entete WHERE Id='{recree}'") == f'{int(version_formulaire) + 1}|{attendu}'
      and nb(f"SELECT COUNT(*) FROM Document_Ligne WHERE EnteteId='{recree}'") == lignes_document
      and 'logo : Logo 2026' in sql(f"SELECT CommentaireVersion FROM Document_Entete WHERE Id='{recree}'"),
      (recree, sql(f"SELECT CONCAT(Version, '|', Nom) FROM Document_Entete WHERE Id='{recree}'"), attendu))
modele_recree = (data(call(SUP, 'GET', f'/rapports/documents/par-document/{recree}')[1]) or {}).get('source', {}).get('modele') or {}
modele_ancien = (data(call(SUP, 'GET', f'/rapports/documents/par-document/{actif}')[1]) or {}).get('source', {}).get('modele') or {}
check('G4 traçabilité : le document recréé a le nouveau modèle, l\'ancien garde le sien ; modèle publié',
      modele_recree.get('versionId', '').lower() == str(version_logo).lower()
      and modele_ancien.get('versionId') == modele_avant.get('versionId')
      and sql(f"SELECT Status FROM ReportTemplateVersion WHERE Id='{version_logo}'") == 'PUBLISHED', (modele_recree, modele_ancien))
st, r = call(GEST, 'POST', f'/rapports/modeles/versions/{version_logo}/changement-logo', {'formulaireIds': [ligne_actuelle]})
check('G5 formulaire archivé : refusé (ses documents gardent leur modèle)', st == 400 and 'archiv' in json.dumps(r, ensure_ascii=False), (st, r))
st, r = call(GEST, 'POST', f'/rapports/modeles/versions/{version_logo}/changement-logo', {'formulaireIds': []})
check('G6 aucun formulaire choisi : refusé', st == 400, (st, r))
nouvelle_ligne = sql(f"SELECT LOWER(CAST(Id AS VARCHAR(36))) FROM Ref_Formulaire WHERE Statut <> 'ARCHIVE' AND VersionGroupId = "
                     f"(SELECT VersionGroupId FROM Ref_Formulaire WHERE Id = '{ligne_actuelle}')")
st, r = call(GEST, 'POST', f'/rapports/modeles/versions/{version_logo}/changement-logo', {'formulaireIds': [nouvelle_ligne]})
check('G6b modèle déjà rattaché à ce formulaire : refusé (envoyer le fichier avec le nouveau logo)',
      st == 400 and 'déjà le modèle' in json.dumps(r, ensure_ascii=False), (st, r))

st, r = call(SUP, 'POST', '/referentiels/formulaires', {'role': 'EN_COURS_DE_ASSEMBLAGE', 'code': 'FE-ASS-PROJET-67', 'natureArticleCode': 'ECROU'})
projet = str(data(r)).lower() if st == 200 else ''
st, r = call(GEST, 'POST', f'/rapports/modeles/versions/{version_logo}/changement-logo', {'formulaireIds': [projet]})
check("G7 modèle dédié à un autre formulaire : refusé, rien n'est modifié",
      st == 400 and "autre formulaire" in json.dumps(r, ensure_ascii=False)
      and nb("SELECT COUNT(*) FROM Ref_Formulaire WHERE CodeReference = 'FE-ASS-PROJET-67'") == 1
      and nb(f"SELECT COUNT(*) FROM ReportTemplateAssignment WHERE FormId='{projet}'") == 0, (st, r))

# ── J. Version de départ donnée à la création du code ; nature des types ; même fichier pour un autre formulaire ──
st, r = call(SUP, 'POST', '/referentiels/formulaires', {'role': 'EN_COURS_DE_ASSEMBLAGE', 'code': 'FE-ASS-V3-67', 'natureArticleCode': 'GOUPILLE', 'version': 3})
form_v3 = str(data(r)).lower() if st == 200 else ''
check('J1 code créé par le superviseur avec sa version de départ (version 3 de la fiche Sopal)',
      st == 200 and sql(f"SELECT Version FROM Ref_Formulaire WHERE Id='{form_v3}'") == '3', (st, r))
st_neg, r_neg = call(SUP, 'POST', '/referentiels/formulaires', {'role': 'EN_COURS_DE_ASSEMBLAGE', 'code': 'FE-ASS-NEG-67', 'natureArticleCode': 'JOINT', 'version': -1})
st_max, r_max = call(SUP, 'POST', '/referentiels/formulaires', {'role': 'EN_COURS_DE_ASSEMBLAGE', 'code': 'FE-ASS-MAX-67', 'natureArticleCode': 'LIM_DEBIT', 'version': 1000})
check('J2 version de départ hors de 0 à 999 : refusée avec un message, aucun code créé',
      (st_neg, st_max) == (400, 400) and 'version de départ' in json.dumps(r_neg, ensure_ascii=False)
      and nb("SELECT COUNT(*) FROM Ref_Formulaire WHERE CodeReference IN ('FE-ASS-NEG-67', 'FE-ASS-MAX-67')") == 0, (st_neg, st_max, r_neg))
st, r = call(SUP, 'POST', '/Document', {'typeDocumentCode': 'PLAN_ASS', 'nom': 'Plan goupille E2E', 'natureArticleCode': 'GOUPILLE',
    'refFormulaireCodeReference': 'FE-ASS-V3-67', 'sections': [{'ordreAffiche': 1, 'libelleSection': 'Contrôle', 'lignes': [
        {'ordreAffiche': 1, 'libelleAffiche': 'Longueur'}]}]})
doc_v3 = (data(r) or {}).get('id') if isinstance(data(r), dict) else None
check('J3 premier document du formulaire : il prend la version de départ (aucune version saisie dans le document)',
      st in (200, 201) and sql(f"SELECT CONCAT(d.Version, '|', f.Version) FROM Document_Entete d JOIN Ref_Formulaire f ON f.Id = d.FormulaireId "
                                 f"WHERE d.Id='{doc_v3}'") == '3|3', (st, r))
st, r = call(SUP, 'GET', '/referentiels/formulaires/roles')
natures = {x.get('code'): x.get('nature') for x in (data(r) or [])}
check("J4 nature des types : résultats d'assemblage et de fabrication = fiches de réponse, plans = documents",
      natures.get('RESULTAT_CONTROLE_ASSEMBLAGE') == 'RESPONSE_SHEET' and natures.get('RESULTAT_CONTROLE_CF') == 'RESPONSE_SHEET'
      and natures.get('EN_COURS_DE_ASSEMBLAGE') == 'DOCUMENT' and natures.get('VERIF_MACHINE') == 'DOCUMENT', natures)
ligne_v3 = sql(f"SELECT LOWER(CAST(FormulaireId AS VARCHAR(36))) FROM Document_Entete WHERE Id='{doc_v3}'")
st, r = envoyer(GEST, 'PLAN_ASS', 'PLAN_ASS.xlsx', MODELE_LIVRE, formulaire=ligne_v3)
dediee = (data(r) or {}).get('version') or {}
check("J5 même fichier Excel pour un autre formulaire : nouveau modèle dédié (V1 de ce formulaire, nommé pour lui), jamais la version d'un autre",
      st == 200 and dediee.get('numero') == 1 and str(dediee.get('id', '')).lower() != str(version1).lower()
      and nb(f"SELECT COUNT(*) FROM ReportTemplateVersion v JOIN Ref_Formulaire f ON f.Id='{ligne_v3}' JOIN ReportTemplate t ON t.Code = v.ReportTemplateCode "
             f"WHERE v.Id='{dediee.get('id')}' AND v.FormGroupId = f.VersionGroupId AND v.Name = f.Designation + NCHAR(32) + NCHAR(8212) + NCHAR(32) + 'FE-ASS-V3-67.3.1'") == 1, (st, r))
st, r = call(GEST, 'POST', f'/rapports/modeles/versions/{version1}/formulaire', {'formulaireId': ligne_v3})
check("J6 rattacher à ce formulaire le modèle d'un autre : refusé (« envoyez le fichier pour … »)",
      st == 400 and 'autre formulaire' in json.dumps(r, ensure_ascii=False), (st, r))

# ── K. Demandes de modèle : automatiques, du superviseur (même version / version suivante), traitées par l'administrateur ──
def demandes(statut=None):
    st, r = call(GEST, 'GET', '/rapports/demandes-modele' + (f'?statut={statut}' if statut else ''))
    return data(r) or [] if st == 200 else []
def demande_de(formulaire_id, statut='PENDING'):
    return next((d for d in demandes(statut) if str(d.get('formulaireId', '')).lower() == str(formulaire_id).lower()), None)
d_auto = demande_de(ligne_v3)
check("K1 nouveau code de formulaire : demande automatique « nouveau formulaire » chez l'administrateur (envoyée par l'application)",
      d_auto and d_auto['origine'] == 'NEW_FORM' and d_auto['demandePar'] is None and d_auto['codeFormulaire'] == 'FE-ASS-V3-67'
      and d_auto['nature'] == 'DOCUMENT', d_auto)
st, r = call(GEST, 'POST', f"/rapports/demandes-modele/{d_auto['id']}/traiter", {'versionId': dediee.get('id')}) if d_auto else (0, {})
check("K2 l'administrateur traite la demande avec le modèle dédié : M1 rattaché, demande traitée, le document déjà créé le reçoit",
      st == 200 and sql(f"SELECT Status FROM ReportTemplateRequest WHERE Id={d_auto['id'] if d_auto else 0}") == 'DONE'
      and (etat(doc_v3).get('source', {}).get('modele') or {}).get('numeroVersion') == 1, (st, r))

st, r = call(SUP, 'POST', f'/rapports/formulaires/{ligne_v3}/demande-modele', {'motif': 'Mise en page', 'changementVersion': 'SAME_VERSION'})
demande_meme = data(r) if st == 200 else None
st2, r2 = call(SUP, 'POST', f'/rapports/formulaires/{ligne_v3}/demande-modele', {'motif': 'Autre', 'changementVersion': 'NEXT_VERSION'})
st3, _ = call(SUP, 'POST', f'/rapports/formulaires/{ligne_v3}/demande-modele', {'motif': '', 'changementVersion': 'SAME_VERSION'})
check('K3 superviseur : demande « même version » envoyée ; une 2e demande en attente, ou sans motif, refusée',
      st == 200 and demande_meme and st2 == 400 and 'attend déjà' in json.dumps(r2, ensure_ascii=False) and st3 == 400, (st, r, st2, r2))
st, r = call(GEST, 'POST', f"/rapports/modeles/versions/{dediee.get('id')}/formulaire", {'formulaireId': ligne_v3})
check("K4 rattachement direct pendant une demande du superviseur : refusé (son choix de version s'applique)",
      st == 400 and 'traitez sa demande' in json.dumps(r, ensure_ascii=False), (st, r))
st, r = envoyer(GEST, 'PLAN_ASS', 'mise_en_page.xlsx', variante(lambda f: f.__setitem__('Z1', 'Mise en page 2')), formulaire=ligne_v3)
m2 = ((data(r) or {}).get('version') or {}).get('id')
st, r = call(GEST, 'POST', f'/rapports/demandes-modele/{demande_meme}/traiter', {'versionId': m2, 'commentaire': 'Marges'})
e_v3 = etat(doc_v3)
check("K5 « même version » traitée : M2 en révision sur la même version (3) ; le document déjà créé garde M1, le modèle actuel est M2",
      st == 200 and sql(f"SELECT CONCAT(Version, '|', Statut) FROM Ref_Formulaire WHERE Id='{ligne_v3}'") == '3|ACTIF'
      and sql(f"SELECT Kind FROM ReportTemplateAssignment WHERE ReportTemplateVersionId='{m2}'") == 'FORM_REVISION'
      and (e_v3['source'].get('modele') or {}).get('numeroVersion') == 1 and (e_v3.get('modeleActuel') or {}).get('numeroVersion') == 2
      and sql(f"SELECT Status FROM ReportTemplateRequest WHERE Id={demande_meme}") == 'DONE', (st, r, e_v3))

st, r = call(SUP, 'POST', f'/rapports/formulaires/{ligne_v3}/demande-modele', {'motif': 'Nouveau logo', 'changementVersion': 'NEXT_VERSION'})
demande_suivante = data(r) if st == 200 else None
st, r = envoyer(GEST, 'PLAN_ASS', 'logo.xlsx', variante(lambda f: f.__setitem__('Z1', 'Logo 2027')), formulaire=ligne_v3)
m3 = ((data(r) or {}).get('version') or {}).get('id')
st, r = call(GEST, 'POST', f'/rapports/demandes-modele/{demande_suivante}/traiter', {'versionId': m3})
ligne_v4 = sql("SELECT LOWER(CAST(Id AS VARCHAR(36))) FROM Ref_Formulaire WHERE CodeReference='FE-ASS-V3-67' AND Statut <> 'ARCHIVE'")
doc_v4 = sql(f"SELECT LOWER(CAST(Id AS VARCHAR(36))) FROM Document_Entete WHERE FormulaireId='{ligne_v4}' AND Statut='ACTIF'")
check("K6 « version suivante » traitée : le formulaire passe en version 4 avec M3, son document est archivé tel quel puis recréé en version 4",
      st == 200 and sql(f"SELECT Version FROM Ref_Formulaire WHERE Id='{ligne_v4}'") == '4'
      and sql(f"SELECT Statut FROM Ref_Formulaire WHERE Id='{ligne_v3}'") == 'ARCHIVE'
      and sql(f"SELECT CONCAT(Version, '|', Statut) FROM Document_Entete WHERE Id='{doc_v3}'") == '3|ARCHIVE' and bool(doc_v4)
      and (etat(doc_v4).get('source', {}).get('modele') or {}).get('numeroVersion') == 3
      and (etat(doc_v3).get('source', {}).get('modele') or {}).get('numeroVersion') == 1
      and sql(f"SELECT CONCAT(r.Status, '|', LOWER(CAST(r.FormId AS VARCHAR(36)))) FROM ReportTemplateRequest r WHERE r.Id={demande_suivante}") == f'DONE|{ligne_v4}', (st, r))

st, r = call(SUP, 'POST', f'/rapports/formulaires/{ligne_v4}/demande-modele', {'motif': 'Erreur', 'changementVersion': 'SAME_VERSION'})
annulee = data(r) if st == 200 else None
st_a, _ = call(SUP, 'POST', f'/rapports/demandes-modele/{annulee}/annuler')
st_b, _ = call(SUP, 'POST', f'/rapports/demandes-modele/{annulee}/annuler')
check('K7 le superviseur annule sa demande (gardée, annulée) ; une demande déjà annulée ne se rannule pas',
      (st_a, st_b) == (200, 400) and sql(f"SELECT Status FROM ReportTemplateRequest WHERE Id={annulee}") == 'CANCELLED', (st_a, st_b))
check("K8 rôles : l'administrateur ne crée pas de demande, le superviseur ne la traite pas (403)",
      call(GEST, 'POST', f'/rapports/formulaires/{ligne_v4}/demande-modele', {'motif': 'x', 'changementVersion': 'SAME_VERSION'})[0] == 403
      and call(SUP, 'POST', f'/rapports/demandes-modele/{annulee}/traiter', {'versionId': m3})[0] == 403)

st, r = call(SUP, 'PUT', f'/referentiels/formulaires/{ligne_v4}/code', {'code': 'FE-ASS-V3B-67'})
ligne_b = sql("SELECT LOWER(CAST(Id AS VARCHAR(36))) FROM Ref_Formulaire WHERE CodeReference='FE-ASS-V3B-67' AND Statut <> 'ARCHIVE'")
doc_b = sql(f"SELECT LOWER(CAST(Id AS VARCHAR(36))) FROM Document_Entete WHERE FormulaireId='{ligne_b}' AND Statut='ACTIF'")
check('K9 changement de code (même version) : la nouvelle ligne reçoit le même modèle M3 (« repris de FE-ASS-V3-67.4 »), son document aussi',
      st == 200 and sql(f"SELECT CONCAT(Version, '|', LOWER(CAST(ReportTemplateVersionId AS VARCHAR(36)))) FROM ReportTemplateAssignment a "
                        f"JOIN Ref_Formulaire f ON f.Id = a.FormId WHERE a.FormId='{ligne_b}'") == f'4|{str(m3).lower()}'
      and 'Repris de FE-ASS-V3-67.4' in sql(f"SELECT Comment FROM ReportTemplateAssignment WHERE FormId='{ligne_b}'")
      and (etat(doc_b).get('source', {}).get('modele') or {}).get('numeroVersion') == 3, (st, r))
historique = sql(f"SELECT STRING_AGG(CONCAT(f.CodeReference, '.', f.Version, '>M', v.VersionNumber), ',') WITHIN GROUP (ORDER BY a.CreationOrder) "
                 f"FROM ReportTemplateAssignment a JOIN Ref_Formulaire f ON f.Id = a.FormId JOIN ReportTemplateVersion v ON v.Id = a.ReportTemplateVersionId "
                 f"WHERE f.VersionGroupId = (SELECT VersionGroupId FROM Ref_Formulaire WHERE Id='{ligne_b}')")
check('K10 relation « code + version → modèle » conservée pour tout le formulaire : FE-ASS-V3-67.3>M1, .3>M2, .4>M3, FE-ASS-V3B-67.4>M3',
      historique == 'FE-ASS-V3-67.3>M1,FE-ASS-V3-67.3>M2,FE-ASS-V3-67.4>M3,FE-ASS-V3B-67.4>M3', historique)

print(f'e2e67: {ok} OK / {ko} KO')
